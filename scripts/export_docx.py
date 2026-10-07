#!/usr/bin/env python3
"""Export the computational working draft from its Markdown source files."""
from pathlib import Path
import re

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = Path(__file__).resolve().parents[1]
PARTS = (
    "chapters/01_results_sequence_quality.md",
    "chapters/02_results_guided_generation.md",
    "chapters/03_methods_computational.md",
    "chapters/04_figure_legends.md",
    "references/method-references.md",
)
PREAMBLE = """# 160 bp 增强子生成：计算部分工作稿

版本日期：2026-10-07。状态：按投稿结构组织的工作稿，尚不能直接投稿。本文档仅覆盖用户确认的计算部分；原始结果、最终模型和图件未核实，待核实项须在投稿前清零。当前采用通用排版，目标期刊尚未指定。
"""


def merged_text():
    return PREAMBLE.rstrip() + "\n\n" + "\n\n".join(
        (ROOT / name).read_text(encoding="utf-8").strip() for name in PARTS
    ) + "\n"


def paragraphs(text):
    for block in re.split(r"\n\s*\n", text.strip()):
        match = re.match(r"^(#{1,6})\s+(.+)$", block, re.DOTALL)
        if match:
            yield len(match.group(1)), match.group(2).strip()
        else:
            yield 0, block.replace("\n", " ").strip()


def export():
    text = merged_text()
    target = ROOT / "manuscript"
    target.mkdir(exist_ok=True)
    (target / "computational_sections_zh.md").write_text(text, encoding="utf-8")
    doc = Document()
    doc.core_properties.title = "160 bp 增强子生成：计算部分工作稿"
    doc.core_properties.author = ""
    doc.core_properties.subject = "Working draft; not submission ready"
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2.5)
    section.left_margin = section.right_margin = Cm(2.5)
    styles = [doc.styles["Normal"]] + [doc.styles[f"Heading {i}"] for i in range(1, 4)]
    for style in styles:
        style.font.name = "Times New Roman"
        style.font.size = Pt(12)
        style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
    doc.styles["Normal"].paragraph_format.line_spacing = 2
    doc.styles["Normal"].paragraph_format.space_after = Pt(6)
    for level, content in paragraphs(text):
        if level:
            doc.add_heading(content, level=min(level, 3))
        else:
            doc.add_paragraph(content)
    numbering = OxmlElement("w:lnNumType")
    numbering.set(qn("w:countBy"), "1")
    numbering.set(qn("w:restart"), "continuous")
    section._sectPr.insert_element_before(numbering, "w:pgNumType", "w:cols", "w:docGrid")
    footer = section.footer.paragraphs[0]
    footer.alignment = 1
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    doc.save(target / "computational_sections_zh.docx")
    print("Exported synchronized Markdown and DOCX working drafts")


if __name__ == "__main__":
    export()
