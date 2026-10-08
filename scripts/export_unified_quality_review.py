"""Export the shared-protocol figure and verified candidate table for review."""
from pathlib import Path
import csv
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1]
doc=Document();section=doc.sections[0];section.page_width=Cm(21);section.page_height=Cm(29.7);section.top_margin=section.bottom_margin=section.left_margin=section.right_margin=Cm(2)
for name in ['Normal','Title','Heading 1']:
 style=doc.styles[name];style.font.name='Times New Roman';style.font.color.rgb=RGBColor(0,0,0);style.font.size=Pt(12 if name=='Normal' else 16 if name=='Title' else 14);style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体' if name=='Normal' else '黑体');style.paragraph_format.line_spacing=1.5
style=doc.styles['Normal'];style.paragraph_format.first_line_indent=Pt(24)
doc.add_heading('DDPM、DiT、UViT：统一协议下的序列质量比较',0)
doc.add_paragraph('2026-10-08审阅版。固定每个候选1000条160bp序列，使用同一真实参照及同一JASPAR2024扫描；指标显示实际权衡，不预设某模型全部最优。')
doc.add_picture(str(ROOT/'figures/model_selection/figure_unified_quality_comparison.png'),width=Cm(17))
doc.add_paragraph('前三行分别是各40批次中6-mer最佳的候选；额外保留DiT1050和原UViT位置WD候选。新增motif命中密度，以接近真实参照38.295次/序列为目标，而非越多越好。新WD为本次统一扫描重算值，与旧图不同版本或不同参照的WD不可直接混用。')
doc.add_page_break();doc.add_heading('核验结果与可支持的结论',1)
rows=list(csv.DictReader((ROOT/'data/verified/unified_quality_comparison.csv').open()))
selected=[('DDPM',1800),('DiT',1200),('DiT',1050),('UViT-v2 config4',1900),('UViT-v2 config4',1800)]
table=doc.add_table(rows=1,cols=5);table.style='Table Grid'
for cell,text in zip(table.rows[0].cells,['模型／epoch','6-mer PCC ↑','GC WD ↓','motif位置WD ↓','命中/序列']):cell.text=text
for model,epoch in selected:
 r=next(r for r in rows if r['model']==model and int(r['epoch'])==epoch)
 for cell,text in zip(table.add_row().cells,[f'{model} / {epoch}',f"{float(r['sixmer_pcc']):.6f}",f"{float(r['gc_wasserstein']):.6f}",f"{float(r['jaspar2024_position_wd']):.6f}",f"{float(r['jaspar2024_hits_per_sequence']):.3f}"]):cell.text=text
doc.add_paragraph('DDPM1800与DiT1200的6-mer PCC接近（差值约0.000119，未证明显著差异）；DiT的GC距离略低，而DDPM的motif位置距离更低。相较本轮UViT-v2 config4候选，DDPM的6-mer及GC更接近参照，但UViT的位置WD更低。选择DDPM的理由可表述为兼顾局部序列组成与motif位置分布，不应写成所有指标最佳。')
doc.add_paragraph('加入数量后，DiT1050为36.896次/序列，最接近真实参照38.295；DDPM1800为33.404。6-mer最高仅适合初筛，未对全部120个checkpoint完成新协议的四指标联合扫描，不能将当前候选称为综合最优。最终应先确定主指标、质量门槛与同等候选筛选规则，避免只比较对DDPM有利的批次。')
doc.add_paragraph('真实参照11984条，来自NG训练数据来源，不是独立测试集。UViT从9000条中按固定seed42无放回抽取1000条；该固定样本用于全部指标。未重新训练或生成，也没有多种子重复，因此结果仅代表现存采样的回顾性比较。')
doc.add_paragraph('FIMO5.4.1、JASPAR2024、p值1e-4、NRDB背景、双链；位置计数不加伪计数。所有候选共享相同参数。原始序列及成员哈希、协议、扫描计数和代码已归档。')
doc.add_paragraph('10月1日引导实验实际使用DDPM2000，不能改标为1800。最终主指标、主权重及是否统一两节论文的权重仍需作者确定。')
target=ROOT/'results/00_model_selection/unified_quality_review_zh.docx';doc.save(target);print('Exported unified quality review DOCX')
